import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from typing import Optional

import httpx

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import freellm  # noqa: E402
from freellm import FreeLLM, AllModelsFailed, tool  # noqa: E402

HOST_TO_GW = {"api.groq.com": "groq", "api.cerebras.ai": "cerebras", "openrouter.ai": "openrouter"}


def ok(text="hi", tool_calls=None, headers=None):
    msg = {"role": "assistant", "content": text}
    if tool_calls:
        msg["tool_calls"] = tool_calls
    return httpx.Response(200, json={"choices": [{"message": msg}], "usage": {"total_tokens": 5}},
                          headers=headers or {})


class Env:
    """Fake clock + scripted gateways."""

    def __init__(self):
        self.t = 1_800_000_000.0  # fixed 'now'
        self.calls: list[tuple[str, str]] = []
        self.script: dict[str, list] = {}   # gateway -> list of responses/callables (last one repeats)

    def handler(self, request: httpx.Request) -> httpx.Response:
        gw = HOST_TO_GW[request.url.host]
        model = json.loads(request.content)["model"]
        self.calls.append((gw, model))
        queue = self.script.get(gw, [ok()])
        item = queue.pop(0) if len(queue) > 1 else queue[0]
        return item(request) if callable(item) else item

    def llm(self, registry=None, **kw) -> FreeLLM:
        d = tempfile.mkdtemp()
        return FreeLLM(registry=registry or make_registry(), state_path=f"{d}/state.json",
                       clock=lambda: self.t, transport=httpx.MockTransport(self.handler), **kw)


def make_registry():
    gws = freellm.all_gateways()
    rows = [
        ("groq", {"id": "llama-3.3-70b-versatile", "context_window": 131072}),
        ("groq", {"id": "llama-3.1-8b-instant", "context_window": 131072}),
        ("groq", {"id": "qwen-2.5-coder-32b", "context_window": 32768}),
        ("cerebras", {"id": "gpt-oss-120b", "context_window": 131072}),
        ("openrouter", {"id": "meta/some-8b:free", "context_length": 4096,
                        "pricing": {"prompt": "0", "completion": "0"}}),
        ("openrouter", {"id": "vendor/paid-70b", "context_length": 65536,
                        "pricing": {"prompt": "0.001", "completion": "0.002"}}),
        ("groq", {"id": "whisper-large-v3"}),
    ]
    models = [m for g, r in rows if (m := freellm.normalize_row(g, gws[g], r))]
    return {"schema_version": 2, "gateways": freellm.public_gateways(gws), "models": models}


class TestEnrichment(unittest.TestCase):
    def test_size_parsing(self):
        p = freellm.parse_size_b
        self.assertEqual(p("llama-3.3-70b-versatile"), 70)
        self.assertEqual(p("mixtral-8x7b-instruct"), 56)
        self.assertEqual(p("qwen3-30b-a3b"), 30)
        self.assertEqual(p("gpt-oss-120b"), 120)
        self.assertEqual(p("gemma-3-27b-it:free"), 27)
        self.assertEqual(p("Llama-3.1-8B-Instant"), 8)
        self.assertIsNone(p("gemini-2.5-flash"))
        self.assertIsNone(p("gemma-3n-e4b-it"))

    def test_classes(self):
        self.assertEqual(freellm.size_class(7), "small")
        self.assertEqual(freellm.size_class(70), "large")
        self.assertEqual(freellm.size_class(None), "unknown")
        self.assertEqual(freellm.context_class(4096), "short")
        self.assertEqual(freellm.context_class(131072), "long")
        self.assertEqual(freellm.context_class(1_000_000), "ultra")

    def test_free_detection_and_filtering(self):
        by = {m["uid"]: m for m in make_registry()["models"]}
        self.assertTrue(by["openrouter/meta/some-8b:free"]["is_free"])
        self.assertFalse(by["openrouter/vendor/paid-70b"]["is_free"])
        self.assertTrue(by["groq/llama-3.3-70b-versatile"]["is_free"])          # free_tier gateway
        self.assertNotIn("groq/whisper-large-v3", by)                           # non-chat dropped
        m = by["groq/llama-3.3-70b-versatile"]
        self.assertEqual((m["size_label"], m["size_class"], m["context_class"], m["long_context"]),
                         ("70B", "large", "long", True))
        self.assertEqual(by["openrouter/meta/some-8b:free"]["context_class"], "short")
        self.assertIn("coding", by["groq/qwen-2.5-coder-32b"]["categories"])

    def test_gemini_prefix_stripped(self):
        m = freellm.normalize_row("gemini", freellm.GATEWAYS["gemini"], {"id": "models/gemini-2.5-flash"})
        self.assertEqual(m["model_id"], "gemini-2.5-flash")
        self.assertEqual(m["context_source"], "hint")


class TestRouting(unittest.TestCase):
    def setUp(self):
        os.environ.update(GROQ_API_KEY="k", CEREBRAS_API_KEY="k", OPENROUTER_API_KEY="k")
        self.env = Env()

    def test_task_selection(self):
        llm = self.env.llm()
        self.assertEqual(llm.candidates("coding")[0]["model_id"], "qwen-2.5-coder-32b")
        self.assertEqual(llm.candidates("fast")[0]["model_id"], "llama-3.1-8b-instant")
        long_doc = llm.candidates("long_doc")
        self.assertTrue(all(m["context_length"] >= 64000 for m in long_doc))
        self.assertNotIn("openrouter/vendor/paid-70b", [m["uid"] for m in llm.candidates("chat")])  # paid excluded

    def test_missing_key_skips_gateway(self):
        del os.environ["CEREBRAS_API_KEY"]
        try:
            llm = self.env.llm()
            self.assertNotIn("cerebras", {m["gateway"] for m in llm.candidates("chat")})
        finally:
            os.environ["CEREBRAS_API_KEY"] = "k"

    def test_daily_429_exhausts_and_skips_without_calling(self):
        # openrouter's daily limit is gateway-wide: one 429 must park the whole gateway
        self.env.script["openrouter"] = [httpx.Response(429, text="free-models-per-day limit exceeded")]
        llm = self.env.llm(registry=only(make_registry(), "openrouter", "cerebras"))
        res = llm.chat("hi", task="fast")                      # openrouter 8B ranks first for 'fast'
        self.assertEqual([g for g, _ in self.env.calls], ["openrouter", "cerebras"])
        self.assertEqual(res.gateway, "cerebras")
        self.assertEqual(llm.gateway_status()["openrouter"]["rate_limit_exhausted_today"], "yes")
        self.env.calls.clear()
        llm.chat("again", task="fast")
        self.assertEqual([g for g, _ in self.env.calls], ["cerebras"])   # zero calls wasted on openrouter

    def test_fallback_and_gateway_yes_flag(self):
        # groq daily-limited on every model -> router moves on to cerebras, status says yes for groq
        self.env.script["groq"] = [httpx.Response(429, text="Rate limit reached ... requests per day (RPD)",
                                                  headers={"retry-after": "7200"})]
        llm = self.env.llm()
        # task="fast" ranks groq's 8B first, so groq is tried first and must fail over
        res = llm.chat("hi", task="fast", gateways=["groq", "cerebras"])
        self.assertEqual(res.gateway, "cerebras")
        self.assertTrue(any(u.startswith("groq/") for u, _ in res.attempts))
        # only the 8B is known-limited so far: lazy discovery, status still "no"
        self.assertEqual(llm.gateway_status()["groq"]["rate_limit_exhausted_today"], "no")
        with self.assertRaises(AllModelsFailed):          # force-try the rest of groq
            llm.chat("hi", gateways=["groq"])
        self.assertEqual(llm.gateway_status()["groq"]["rate_limit_exhausted_today"], "yes")
        self.assertEqual(llm.gateway_status()["cerebras"]["rate_limit_exhausted_today"], "no")
        self.env.calls.clear()
        llm.chat("hi")
        self.assertNotIn("groq", [g for g, _ in self.env.calls])

    def test_minute_429_is_short_cooldown(self):
        self.env.script["groq"] = [httpx.Response(429, text="rate limit", headers={"retry-after": "10"}), ok()]
        llm = self.env.llm(registry=only(make_registry(), "groq"))
        first = llm.candidates("fast")[0]["uid"]
        res = llm.chat("hi", task="fast")
        self.assertNotEqual(f"groq/{res.model_id}", first)          # moved to another groq model
        self.assertEqual(llm.gateway_status()["groq"]["rate_limit_exhausted_today"], "no")
        self.assertNotIn(first, [m["uid"] for m in llm.candidates("fast")])
        self.env.t += 11
        self.assertIn(first, [m["uid"] for m in llm.candidates("fast")])  # cooldown expired

    def test_404_disables_model_for_a_day(self):
        self.env.script["groq"] = [httpx.Response(404, text="model not found"), ok()]
        llm = self.env.llm(registry=only(make_registry(), "groq"))
        first = llm.candidates("chat")[0]["uid"]
        llm.chat("hi")
        self.env.t += 3600 * 5
        self.assertNotIn(first, [m["uid"] for m in llm.candidates("chat")])

    def test_header_remaining_zero_marks_exhausted(self):
        self.env.script["groq"] = [ok(headers={"x-ratelimit-remaining-requests": "0",
                                               "x-ratelimit-reset-requests": "3h"})]
        llm = self.env.llm(registry=only(make_registry(), "groq"))
        first = llm.candidates("chat")[0]["uid"]
        llm.chat("hi")
        self.assertNotIn(first, [m["uid"] for m in llm.candidates("chat")])

    def test_day_rollover_resets(self):
        self.env.script["groq"] = [httpx.Response(402, text="payment required")]
        llm = self.env.llm(registry=only(make_registry(), "groq"))
        with self.assertRaises(AllModelsFailed):
            llm.chat("hi")
        self.assertEqual(llm.gateway_status()["groq"]["rate_limit_exhausted_today"], "yes")
        self.env.t += 86400 + 5
        self.assertEqual(llm.gateway_status()["groq"]["rate_limit_exhausted_today"], "no")
        self.assertTrue(llm.candidates("chat"))

    def test_state_persists_between_instances(self):
        self.env.script["groq"] = [httpx.Response(402, text="payment required")]
        d = tempfile.mkdtemp()
        mk = lambda: FreeLLM(registry=only(make_registry(), "groq"), state_path=f"{d}/s.json",
                             clock=lambda: self.env.t, transport=httpx.MockTransport(self.env.handler))
        a = mk()
        with self.assertRaises(AllModelsFailed):
            a.chat("hi")
        self.assertEqual(mk().candidates("chat"), [])

    def test_all_fail_raises_with_attempts(self):
        self.env.script["groq"] = [httpx.Response(500, text="boom")]
        llm = self.env.llm(registry=only(make_registry(), "groq"))
        with self.assertRaises(AllModelsFailed) as cm:
            llm.chat("hi")
        self.assertTrue(cm.exception.attempts)


def only(reg, *gateways):
    return {**reg, "models": [m for m in reg["models"] if m["gateway"] in gateways]}


class TestTools(unittest.TestCase):
    def setUp(self):
        os.environ.update(GROQ_API_KEY="k", CEREBRAS_API_KEY="k", OPENROUTER_API_KEY="k")
        self.env = Env()

    def test_schema_generation(self):
        @tool
        def get_gst_rate(hsn_code: str, intra_state: bool = True, qty: Optional[int] = None):
            """Look up the GST rate for an HSN code.

            Args:
                hsn_code: 4 to 8 digit HSN code
                intra_state: true for CGST+SGST
            """
        p = get_gst_rate.schema["function"]
        self.assertEqual(p["name"], "get_gst_rate")
        self.assertEqual(p["description"], "Look up the GST rate for an HSN code.")
        self.assertEqual(p["parameters"]["required"], ["hsn_code"])
        self.assertEqual(p["parameters"]["properties"]["intra_state"]["type"], "boolean")
        self.assertEqual(p["parameters"]["properties"]["qty"]["type"], "integer")
        self.assertEqual(p["parameters"]["properties"]["hsn_code"]["description"], "4 to 8 digit HSN code")

    def test_tool_loop(self):
        seen = {}

        @tool
        def add(a: int, b: int) -> int:
            """Add two numbers."""
            seen["args"] = (a, b)
            return a + b

        def respond(request):
            body = json.loads(request.content)
            if any(m["role"] == "tool" for m in body["messages"]):
                tool_msg = [m for m in body["messages"] if m["role"] == "tool"][0]
                return ok(f"answer={tool_msg['content']}")
            assert body["tools"][0]["function"]["name"] == "add"
            return ok("", tool_calls=[{"id": "c1", "type": "function",
                                       "function": {"name": "add", "arguments": '{"a": 2, "b": 3}'}}])

        self.env.script["groq"] = [respond]
        llm = self.env.llm(registry=only(make_registry(), "groq"))
        res = llm.run("what is 2+3?", [add])
        self.assertEqual(seen["args"], (2, 3))
        self.assertEqual(res.text, "answer=5")
        self.assertEqual(res.steps[0]["tool"], "add")

    def test_tool_error_is_reported_to_model(self):
        @tool
        def boom(x: int) -> int:
            """Fail."""
            raise ValueError("nope")

        def respond(request):
            msgs = json.loads(request.content)["messages"]
            t = [m for m in msgs if m["role"] == "tool"]
            if t:
                return ok("saw:" + t[0]["content"])
            return ok("", tool_calls=[{"id": "c", "type": "function",
                                       "function": {"name": "boom", "arguments": '{"x": 1}'}}])

        self.env.script["groq"] = [respond]
        llm = self.env.llm(registry=only(make_registry(), "groq"))
        self.assertIn("ValueError", llm.run("go", [boom]).text)

    def test_tools_require_tool_capable_models(self):
        reg = make_registry()
        for m in reg["models"]:
            if m["uid"] == "groq/llama-3.1-8b-instant":
                m["capabilities"]["tools"] = False
        llm = self.env.llm(registry=reg)
        self.assertNotIn("groq/llama-3.1-8b-instant", [m["uid"] for m in llm.candidates("agent")])


class TestRefreshMerge(unittest.TestCase):
    def test_merge_keeps_failed_gateway_and_marks_missing(self):
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
        import refresh_registry as rr
        reg = make_registry()["models"]
        now = 1_800_000_000.0
        fresh = {"groq": [m for m in reg if m["gateway"] == "groq" and m["model_id"] != "llama-3.1-8b-instant"]}
        merged = {m["uid"]: m for m in rr.merge(reg, fresh, now)}
        self.assertEqual(merged["groq/llama-3.1-8b-instant"]["status"], "missing")
        self.assertFalse(merged["groq/llama-3.1-8b-instant"]["available"])
        self.assertEqual(merged["cerebras/gpt-oss-120b"]["status"], "active")   # gateway not refreshed -> kept
        later = rr.merge(list(merged.values()), fresh, now + 5 * 86400)
        self.assertNotIn("groq/llama-3.1-8b-instant", [m["uid"] for m in later])  # pruned after 3 days


class TestMistralRows(unittest.TestCase):
    def test_mistral_base_type_is_kept_and_non_chat_dropped(self):
        cfg = freellm.GATEWAYS["mistral"]
        chat = {"id": "mistral-large-latest", "type": "base", "max_context_length": 131072,
                "capabilities": {"completion_chat": True, "function_calling": True, "vision": False}}
        vis = {"id": "pixtral-large-latest", "type": "base", "max_context_length": 131072,
               "capabilities": {"completion_chat": True, "function_calling": True, "vision": True}}
        emb = {"id": "mistral-embed", "type": "base", "capabilities": {"completion_chat": False}}
        ocr = {"id": "mistral-ocr-latest", "type": "base", "capabilities": {"completion_chat": False}}
        m = freellm.normalize_row("mistral", cfg, chat)
        self.assertIsNotNone(m)
        self.assertTrue(m["capabilities"]["tools"])
        self.assertEqual(m["capability_source"], "declared")
        self.assertEqual(m["context_class"], "long")
        self.assertTrue(freellm.normalize_row("mistral", cfg, vis)["capabilities"]["vision"])
        self.assertIsNone(freellm.normalize_row("mistral", cfg, emb))
        self.assertIsNone(freellm.normalize_row("mistral", cfg, ocr))

    def test_together_style_types_still_filtered(self):
        cfg = freellm.GATEWAYS["together"]
        self.assertIsNone(freellm.normalize_row("together", cfg, {"id": "x/y-8b", "type": "embedding"}))
        self.assertIsNotNone(freellm.normalize_row("together", cfg, {"id": "x/y-8b", "type": "chat",
                                                                    "pricing": {"input": 0, "output": 0}}))

    def test_tokenharbor_url(self):
        self.assertEqual(freellm.GATEWAYS["tokenharbor"]["base_url"], "https://tokenharbor.ai/v1")


class TestAutoContext(unittest.TestCase):
    def setUp(self):
        os.environ.update(GROQ_API_KEY="k", CEREBRAS_API_KEY="k", OPENROUTER_API_KEY="k")
        self.env = Env()

    def test_small_prompt_may_use_short_context_model(self):
        llm = self.env.llm()
        ids = [m["model_id"] for m in llm.candidates("fast")]
        self.assertIn("meta/some-8b:free", ids)            # 4K-context model is eligible for small prompts
        self.assertLess(llm.estimate_tokens([{"role": "user", "content": "hi"}]), 2000)

    def test_big_prompt_skips_small_context_models(self):
        llm = self.env.llm()
        big = "x" * 200_000                                  # ~50K tokens
        res = llm.chat(big, task="fast")
        self.assertIn(res.model_id, {"llama-3.3-70b-versatile", "llama-3.1-8b-instant", "gpt-oss-120b"})
        used = {m for _, m in self.env.calls}
        self.assertNotIn("meta/some-8b:free", used)          # 4K
        self.assertNotIn("qwen-2.5-coder-32b", used)         # 32K < ~58K needed

    def test_nothing_fits_gives_clear_error(self):
        llm = self.env.llm()
        with self.assertRaises(AllModelsFailed) as cm:
            llm.chat("x" * 2_000_000)                        # ~500K tokens
        self.assertIn("tokens of context", str(cm.exception))
        self.assertEqual(self.env.calls, [])                 # failed before wasting any request

    def test_can_disable(self):
        llm = self.env.llm()
        res = llm.chat("x" * 200_000, task="fast", auto_context=False)
        self.assertTrue(res.text)

    def test_task_floor_not_lowered_by_explicit_min_context(self):
        llm = self.env.llm()
        ctxs = [m["context_length"] for m in llm.candidates("long_doc", min_context=1000)]
        self.assertTrue(all(c >= 64000 for c in ctxs))


class TestTables(unittest.TestCase):
    def setUp(self):
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
        import render_tables
        self.rt = render_tables

    def test_markdown_and_csv(self):
        reg = freellm.seed_registry()
        md = self.rt.render_markdown(reg)
        for heading in ("Which context size do I need?", "Best picks per task", "grouped by context class"):
            self.assertIn(heading, md)
        self.assertIn("| groq | `llama-3.3-70b-versatile` | 70B | 128K |", md)
        d = Path(tempfile.mkdtemp())
        self.rt.write_tables(reg, d)
        self.assertTrue((d / "models.csv").exists())
        self.assertTrue((d / "models.xlsx").exists())

    def test_pipe_in_names_is_escaped(self):
        self.assertEqual(self.rt.esc("a|b"), "a\\|b")

    def test_fmt_ctx(self):
        f = self.rt.fmt_ctx
        self.assertEqual((f(131072), f(1048576), f(32768), f(None), f(96000)), ("128K", "1M", "32K", "?", "96K"))
        self.assertEqual(f(1048576, "hint"), "1M≈")


class TestHeuristics(unittest.TestCase):
    def test_context_hint_only_for_gemini(self):
        cer = freellm.normalize_row("cerebras", freellm.GATEWAYS["cerebras"], {"id": "gpt-oss-120b"})
        gem = freellm.normalize_row("gemini", freellm.GATEWAYS["gemini"], {"id": "models/gemini-2.5-flash"})
        self.assertIsNone(cer["context_length"])
        self.assertEqual(cer["context_class"], "unknown")
        self.assertEqual(gem["context_length"], 1_048_576)

    def test_family_quality_prior_for_unsized_models(self):
        q = lambda g, i: freellm.normalize_row(g, freellm.GATEWAYS[g], {"id": i})["quality_hint"]
        self.assertGreater(q("gemini", "gemini-2.5-pro"), q("gemini", "gemini-2.5-flash"))
        self.assertGreater(q("gemini", "gemini-2.5-flash"), q("gemini", "gemini-2.5-flash-lite"))
        self.assertGreater(q("mistral", "mistral-large-latest"), q("mistral", "mistral-small-latest"))
        self.assertGreaterEqual(q("gemini", "gemini-2.5-pro"), 85)


class TestDecisionTable(unittest.TestCase):
    def setUp(self):
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
        import render_tables
        self.rt = render_tables
        gws = freellm.all_gateways()
        self.reg = freellm.seed_registry()
        paid = freellm.normalize_row("openrouter", gws["openrouter"], {
            "id": "vendor/paid-70b", "context_length": 65536, "pricing": {"prompt": "0.001", "completion": "0.002"}})
        vis = [("nvidia", {"id": "meta/llama-3.2-11b-vision-instruct"}),
               ("openrouter", {"id": "qwen/qwen2.5-vl-72b-instruct:free", "context_length": 32768,
                               "pricing": {"prompt": "0", "completion": "0"}})]
        self.reg["models"] += [paid] + [freellm.normalize_row(g, gws[g], r) for g, r in vis]
        gone = freellm.normalize_row("groq", gws["groq"], {"id": "removed-8b"})
        gone["status"] = "missing"
        self.reg["models"].append(gone)

    def read_csv(self):
        import csv
        d = Path(tempfile.mkdtemp())
        self.rt.write_tables(self.reg, d)
        with open(d / "models.csv", encoding="utf-8", newline="") as fh:
            return d, list(csv.reader(fh))

    def test_canonical_key_clubs_same_model_only(self):
        k = self.rt.canonical_key
        self.assertEqual(k("openai/gpt-oss-120b:free"), k("gpt-oss-120b"))
        self.assertEqual(k("llama3.1-8b"), k("meta/llama-3.1-8b-instruct"))
        self.assertEqual(k("llama-3.1-8b-instant"), k("llama3.1-8b"))
        self.assertEqual(k("qwen-3-32b"), k("qwen/qwen3-32b"))
        self.assertEqual(k("llama-3.3-70b-versatile"), k("meta-llama/llama-3.3-70b-instruct:free"))
        self.assertNotEqual(k("gemini-2.5-flash"), k("gemini-2.5-flash-lite"))
        self.assertNotEqual(k("llama-3.3-70b-instruct"), k("llama-3.1-70b-instruct"))
        self.assertNotEqual(k("gpt-oss-20b"), k("gpt-oss-120b"))

    def test_free_only_grouped_and_ordered(self):
        _, rows = self.read_csv()
        header, data = rows[0], rows[1:]
        self.assertEqual(header, self.rt.COLUMNS)
        main = []
        for r in data:
            if not r[0]:
                break
            main.append(r)
        ids = [r[6] for r in main]
        self.assertNotIn("vendor/paid-70b", " ".join(ids))       # paid excluded
        self.assertNotIn("removed-8b", " ".join(ids))            # missing excluded
        ranks = [int(r[0]) for r in main]
        self.assertEqual(ranks, sorted(ranks))                   # groups contiguous, numbered in order
        tiers = [int(r[3]) for r in main]
        self.assertEqual(tiers, sorted(tiers, reverse=True))     # highest intelligence first
        gpt = [r for r in main if r[1] == "gpt-oss-120b"]
        self.assertEqual({r[5] for r in gpt}, {"groq", "cerebras"})   # same model, two gateways, one group
        self.assertEqual(len({r[0] for r in gpt}), 1)
        self.assertEqual(gpt[0][17], "2")
        llama = [r for r in main if r[0] == [x for x in main if x[1].startswith("llama-3.3-70b")][0][0]]
        self.assertEqual({r[5] for r in llama}, {"groq", "openrouter", "nvidia"})
        ctxs = [int(r[8]) if r[8] else 0 for r in llama]
        self.assertEqual(ctxs, sorted(ctxs, reverse=True))       # largest context first inside a group

    def test_labels_and_emoji(self):
        _, rows = self.read_csv()
        top = [r for r in rows[1:] if r[1] == "gpt-oss-120b"][0]
        self.assertEqual(top[2], "🟣 Top")
        self.assertTrue(top[7].startswith("🟢 Long"))
        unknown = [r for r in rows[1:] if r[5] == "cerebras"][0]
        self.assertTrue(unknown[7].startswith("⚪ Unknown"))
        self.assertTrue(any(r[7].startswith("🔵 Ultra") for r in rows[1:]))

    def test_vision_specialists_below_dual_stay_on_top(self):
        _, rows = self.read_csv()
        banner = [i for i, r in enumerate(rows) if r[0].startswith("VISION-SPECIALIST")]
        self.assertEqual(len(banner), 1)
        b = banner[0]
        self.assertEqual(rows[b - 1], [""] * len(self.rt.COLUMNS))
        self.assertEqual(rows[b + 1], self.rt.COLUMNS)
        above = [r[1] for r in rows[1:b - 1]]
        below = [r[1] for r in rows[b + 2:]]
        self.assertIn("gemini-2.5-pro", above)                    # dual (text + vision) stays on top
        self.assertIn("gemma-3-27b-it", above)
        self.assertTrue(any(r[9] == "Text + Vision" for r in rows[1:b - 1]))
        self.assertEqual(set(below), {"qwen2.5-vl-72b-instruct", "llama-3.2-11b-vision-instruct"})
        self.assertTrue(all(r[9] == "Vision specialist" for r in rows[b + 2:]))

    def test_vision_specialist_detection(self):
        v = self.rt.is_vision_specialist
        for name in ("qwen/qwen2.5-vl-72b-instruct", "meta/llama-3.2-11b-vision-instruct", "llava-1.6", "nvidia/neva-22b"):
            self.assertTrue(v(name), name)
        for name in ("gemini-2.5-pro", "google/gemma-3-27b-it", "mistral-small-latest", "pixtral-large-latest",
                     "llama-3.3-70b-versatile"):
            self.assertFalse(v(name), name)

    def test_xlsx_colours_filter_and_freeze(self):
        from openpyxl import load_workbook
        d, rows = self.read_csv()
        wb = load_workbook(d / "models.xlsx")
        ws = wb["Free models"]
        self.assertIn("How to read", wb.sheetnames)
        n_main = next(i for i, r in enumerate(rows) if not r[0]) - 1          # data rows above the blank row
        self.assertEqual(ws.auto_filter.ref, f"A1:R{n_main + 1}")             # filter covers the main table only
        self.assertEqual(ws.freeze_panes, "C2")
        col = {c.value: c.column for c in ws[1]}
        fill = lambda r, name: ws.cell(r, col[name]).fill.start_color.rgb[-6:]
        self.assertEqual(ws.cell(2, col["Intelligence"]).value, "Top")
        self.assertEqual(fill(2, "Intelligence"), "D5B8EA")                   # Top = purple
        self.assertEqual(fill(2, "Context class"), "C6EFCE")                  # Long = green
        gem = next(r for r in range(2, n_main + 2) if ws.cell(r, col["Model"]).value == "gemini-2.5-pro")
        self.assertEqual(fill(gem, "Context class"), "A9D6F5")                # Ultra = blue
        unk = next(r for r in range(2, n_main + 2) if ws.cell(r, col["Gateway"]).value == "cerebras")
        self.assertEqual(fill(unk, "Context class"), "D9D9D9")                # Unknown = grey
        banner_row = n_main + 3
        self.assertTrue(str(ws.cell(banner_row, 1).value).startswith("VISION-SPECIALIST"))

    def test_empty_registry_does_not_crash(self):
        d = Path(tempfile.mkdtemp())
        self.rt.write_tables({"models": [], "updated_at": "x", "source": "t"}, d)
        self.assertEqual(len((d / "models.csv").read_text().splitlines()), 1)


class TestIntelligence(unittest.TestCase):
    def test_tiers_follow_size(self):
        def tier(g, i):
            return freellm.normalize_row(g, freellm.GATEWAYS[g], {"id": i})["intelligence_tier"]
        self.assertEqual(tier("groq", "tiny-1b-instruct"), 1)
        self.assertEqual(tier("groq", "llama-3.1-8b-instant"), 2)
        self.assertEqual(tier("groq", "gemma-27b"), 3)
        self.assertEqual(tier("groq", "llama-3.3-70b-versatile"), 4)
        self.assertEqual(tier("groq", "huge-405b"), 5)
        self.assertEqual(tier("gemini", "gemini-2.5-pro"), 5)
        self.assertEqual(freellm.normalize_row("groq", freellm.GATEWAYS["groq"], {"id": "llama-3.3-70b-versatile"})
                         ["intelligence_label"], "Strong")

    def test_router_filter_and_table_column(self):
        os.environ.update(GROQ_API_KEY="k", CEREBRAS_API_KEY="k", OPENROUTER_API_KEY="k")
        llm = Env().llm()
        strong = llm.candidates("chat", min_intelligence=4)
        self.assertTrue(strong)
        self.assertTrue(all(m["intelligence_tier"] >= 4 for m in strong))
        self.assertLess(len(strong), len(llm.candidates("chat")))
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
        import render_tables
        md = render_tables.render_markdown(freellm.seed_registry())
        self.assertIn("| Intelligence |", md)
        self.assertIn("★★★★☆ Strong (86)", md)
        old_row = {"quality_hint": 86}                      # registry written before this field existed
        self.assertEqual(render_tables.tier_of(old_row), 4)


class TestRefreshLog(unittest.TestCase):
    def test_record_log_and_table(self):
        sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
        import refresh_registry as rr
        reg = make_registry()["models"]
        groq = [m for m in reg if m["gateway"] == "groq"]
        results = [
            {"gateway": "groq", "models": groq[:-1], "listed": 9, "http_status": 200, "latency_ms": 120,
             "status": "ok", "error": None},
            {"gateway": "gemini", "models": None, "listed": 0, "http_status": 401, "latency_ms": 80,
             "status": "error", "error": "401 Unauthorized"},
            {"gateway": "mistral", "models": None, "listed": 0, "http_status": None, "latency_ms": None,
             "status": "skipped", "error": "API key not configured"},
        ]
        rec = rr.build_record(results, reg, 1_800_000_000.0)
        self.assertEqual((rec["ok"], rec["failed"], rec["skipped"]), (1, 1, 1))
        self.assertEqual(rec["gateways"]["groq"]["removed_count"], 1)
        self.assertEqual(rec["gateways"]["groq"]["added_count"], 0)
        self.assertEqual(rec["gateways"]["gemini"]["http_status"], 401)
        self.assertIn("gemini", rr.render_table(rec))
        self.assertIn("| groq |", rr.render_table(rec, markdown=True))
        d = tempfile.mkdtemp()
        log = Path(d) / "log.jsonl"
        for _ in range(5):
            rr.append_log(log, rec, keep=3)
        lines = log.read_text().splitlines()
        self.assertEqual(len(lines), 3)
        self.assertEqual(json.loads(lines[-1])["failed"], 1)


if __name__ == "__main__":
    unittest.main()
